"""学习资料附件：上传/转写/归属/计划与聊天链路。"""
from __future__ import annotations

import io

import pytest

from app import attachments as attachments_svc
from app import store
from app.schemas import Plan, PlanSource

from conftest import auth_headers


@pytest.fixture
def no_auto_transcribe(monkeypatch):
    """上传端点不真起后台转写线程（测试里手动同步调用）。"""
    monkeypatch.setattr(attachments_svc, "start_transcription", lambda uid, aid: None)


def _upload(client, filename: str, content: bytes, mime: str = "application/octet-stream"):
    return client.post(
        "/api/files",
        files={"file": (filename, io.BytesIO(content), mime)},
    )


# ---------------------------------------------------------------------------
# 上传与归属
# ---------------------------------------------------------------------------

def test_upload_txt_and_transcribe(client, no_auto_transcribe, admin_user):
    r = _upload(client, "资料.txt", "Python 装饰器是…".encode("utf-8"), "text/plain")
    assert r.status_code == 200, r.text
    att = r.json()
    assert att["mime"] == "text/plain"
    assert att["transcriptStatus"] == "pending"

    attachments_svc.transcribe_attachment(str(admin_user["id"]), att["id"])
    r2 = client.get(f"/api/files/{att['id']}")
    assert r2.status_code == 200
    body = r2.json()
    assert body["transcriptStatus"] == "done"
    assert "装饰器" in body["transcript"]


def test_upload_rejects_unsupported_type(client, no_auto_transcribe):
    r = _upload(client, "evil.exe", b"MZ...", "application/octet-stream")
    assert r.status_code == 415


def test_upload_size_limit(client, no_auto_transcribe, monkeypatch):
    monkeypatch.setattr(attachments_svc, "MAX_UPLOAD_BYTES", 10)
    r = _upload(client, "big.txt", b"x" * 11, "text/plain")
    assert r.status_code == 413


def test_attachment_owner_isolation(client, no_auto_transcribe, admin_user, normal_user):
    r = _upload(client, "mine.txt", b"hello", "text/plain")
    aid = r.json()["id"]
    # 其他用户看不到: 404
    r2 = client.get(f"/api/files/{aid}", headers=auth_headers(normal_user))
    assert r2.status_code == 404
    r3 = client.get(f"/api/files/{aid}/raw", headers=auth_headers(normal_user))
    assert r3.status_code == 404
    # 删除后本人也 404
    assert client.delete(f"/api/files/{aid}").status_code == 200
    assert client.get(f"/api/files/{aid}").status_code == 404


def test_raw_returns_original_bytes(client, no_auto_transcribe):
    r = _upload(client, "notes.txt", b"raw-body-check", "text/plain")
    aid = r.json()["id"]
    raw = client.get(f"/api/files/{aid}/raw")
    assert raw.status_code == 200
    assert raw.content == b"raw-body-check"


def test_retry_transcribe(client, no_auto_transcribe, admin_user):
    r = _upload(client, "a.txt", b"content", "text/plain")
    aid = r.json()["id"]
    r2 = client.post(f"/api/files/{aid}/transcribe")
    assert r2.status_code == 200


# ---------------------------------------------------------------------------
# 转写派发: docx 抽文本 / 不支持类型 / 空文件
# ---------------------------------------------------------------------------

def _make_docx(text: str) -> bytes:
    from docx import Document as DocxDocument

    doc = DocxDocument()
    doc.add_heading("一级标题", level=1)
    doc.add_paragraph(text)
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "甲"
    table.rows[0].cells[1].text = "乙"
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def test_transcribe_docx_extracts_text(client, no_auto_transcribe, admin_user):
    r = _upload(client, "讲义.docx", _make_docx("线性代数基础内容"), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")
    assert r.status_code == 200
    aid = r.json()["id"]
    attachments_svc.transcribe_attachment(str(admin_user["id"]), aid)
    body = client.get(f"/api/files/{aid}").json()
    assert body["transcriptStatus"] == "done"
    assert "线性代数" in body["transcript"]
    assert "| 甲 | 乙 |" in body["transcript"]


def test_transcribe_persists_failure_message(client, no_auto_transcribe, admin_user, monkeypatch):
    # 视觉模型未配置的图片 → failed + 可读原因
    r = _upload(client, "photo.png", b"\x89PNG fake", "image/png")
    aid = r.json()["id"]
    attachments_svc.transcribe_attachment(str(admin_user["id"]), aid)
    body = client.get(f"/api/files/{aid}").json()
    assert body["transcriptStatus"] == "failed"
    assert body["transcript"]  # 失败原因写进 transcript 供前端展示


def test_sniff_mime_by_extension():
    assert attachments_svc.sniff_mime("a.PDF", "text/html") == "application/pdf"
    assert attachments_svc.sniff_mime("b.DOCX", "") .startswith("application/vnd.openxmlformats")
    assert attachments_svc.sniff_mime("c.mp4", "video/mp4") == ""  # 音视频 MVP 不做
    assert attachments_svc.sniff_mime("noext", "text/plain") == "text/plain"


# ---------------------------------------------------------------------------
# 计划创建 / 教练聊天 挂附件
# ---------------------------------------------------------------------------

@pytest.fixture
def ready_attachment(admin_user):
    """一个转写完成的附件（直接入库，不走上传/转写）。"""
    row = store.create_attachment(
        str(admin_user["id"]), "讲义.pdf", "application/pdf", b"%PDF-fake"
    )
    store.update_transcript(str(row["id"]), "这是讲义正文：函数式编程入门。", "done")
    return str(row["id"])


def test_create_plan_with_attachments(client, fake_coach, ready_attachment, admin_user):
    from _factories import make_plan

    fake_coach.plan = make_plan(modules=2)
    r = client.post("/api/plans", json={
        "input": "", "mode": "topic", "attachmentIds": [ready_attachment],
    })
    assert r.status_code == 200, r.text
    src = fake_coach.last_source
    assert src.mode == "materials"  # 带附件强制走资料模式
    assert "函数式编程入门" in src.input
    assert src.attachmentIds == [ready_attachment]


def test_create_plan_409_when_transcript_not_ready(client, fake_coach, admin_user):
    from _factories import make_plan

    fake_coach.plan = make_plan(modules=1)
    row = store.create_attachment(str(admin_user["id"]), "慢.pdf", "application/pdf", b"%PDF")
    r = client.post("/api/plans", json={"input": "", "mode": "topic", "attachmentIds": [str(row["id"])]})
    assert r.status_code == 409


def test_coach_stream_with_attachments(client, fake_coach, ready_attachment):
    r = client.post("/api/coach/stream", json={
        "goal": "帮我总结这份资料",
        "history": [],
        "attachmentIds": [ready_attachment],
    })
    assert r.status_code == 200
    assert "函数式编程入门" in fake_coach.last_materials
